"""
Script to generate a beautifully styled Microsoft Word (.docx) report
for the Agentic Text-to-SQL & Document RAG Engine project.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding inside table cells (in twentieths of a point / dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, color="D3D3D3", sz="4", val="single"):
    """Sets subtle light borders on a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def add_callout(doc, text, bold_prefix="Key Takeaway: "):
    """Creates a stylized quote/callout block with left border and soft background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    # Left accent border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run_bold = p.add_run(bold_prefix)
    run_bold.bold = True
    run_bold.font.name = "Calibri"
    run_bold.font.size = Pt(10.5)
    run_bold.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10.5)
    run_text.font.italic = True
    run_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # Empty paragraph after table for spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def format_heading_1(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title_text)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)  # Deep Navy
    return p

def format_heading_2(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title_text)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(0x2E, 0x6B, 0x9E)  # Steel Blue
    return p

def format_heading_3(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title_text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return p

def add_body_paragraph(doc, text, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_bullet(doc, bold_lead, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    r_lead = p.add_run(bold_lead + " ")
    r_lead.bold = True
    r_lead.font.name = "Calibri"
    r_lead.font.size = Pt(10.5)
    r_lead.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def create_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Format Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Format Data Rows
    for r_idx, row_values in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
            set_cell_borders(row_cells[c_idx], color="E5E7EB", sz="4")
            p = row_cells[c_idx].paragraphs[0]
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Set Column Widths if provided
    if col_widths:
        for row in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)

def add_figure(doc, image_filename, caption_text, width_inches=5.8):
    """Embeds an image from the screenshots/ directory with centered alignment and italicized caption."""
    img_path = os.path.join(os.path.dirname(__file__), "screenshots", image_filename)
    if not os.path.isfile(img_path):
        img_path = os.path.join(os.path.dirname(__file__), image_filename)
    
    if os.path.isfile(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        p_cap.paragraph_format.keep_with_next = False
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def generate_report():
    doc = Document()

    # 1 Inch Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ---------------------------------------------------------
    # COVER / HEADER BLOCK
    # ---------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(4)
    r_title = title_p.add_run("Project Technical Report")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(26)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    r_sub = sub_p.add_run("Autonomous Multi-Agent Text-to-SQL & Document Retrieval-Augmented Generation (RAG) Engine")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x2E, 0x6B, 0x9E)

    # Metadata Banner Box
    meta_tbl = doc.add_table(rows=1, cols=4)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    
    labels = [
        ("PROJECT DOMAIN", "GenAI & Enterprise Data"),
        ("DATE", "September 2026"),
        ("STATUS", "Production Specification"),
        ("ARCHITECTURE", "LangGraph Supervisor")
    ]
    
    for i, (lbl, val) in enumerate(labels):
        c = meta_tbl.cell(0, i)
        c.width = Inches(1.625)
        set_cell_background(c, "F0F4F8")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        set_cell_borders(c, color="CBD5E1", sz="4")
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        
        r1 = p.add_run(f"{lbl}\n")
        r1.font.name = "Calibri"
        r1.font.size = Pt(7.5)
        r1.bold = True
        r1.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        r2 = p.add_run(val)
        r2.font.name = "Calibri"
        r2.font.size = Pt(9.5)
        r2.bold = True
        r2.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(12)
    p_div.paragraph_format.space_after = Pt(16)

    # ---------------------------------------------------------
    # 1. PROJECT TITLE & PROBLEM STATEMENT
    # ---------------------------------------------------------
    format_heading_1(doc, "1. Project Title & Problem Statement")
    
    format_heading_2(doc, "1.1 Project Title")
    add_body_paragraph(
        doc,
        "Agentic Text-to-SQL & Document RAG Engine: An Autonomous Dual-Path Generative AI Architecture "
        "for Unified Structured and Unstructured Enterprise Knowledge Retrieval."
    )

    format_heading_2(doc, "1.2 Context & Industry Background")
    add_body_paragraph(
        doc,
        "Every modern digital enterprise and e-commerce platform relies on two fundamentally disparate "
        "categories of corporate knowledge to drive day-to-day operations and strategic decisions:"
    )
    add_bullet(
        doc,
        "1. Structured Operational Data:",
        "Millions of transactional rows stored in relational database tables (e.g., PostgreSQL). "
        "This includes customer accounts, order transaction histories, line-item pricing, product catalogs, "
        "delivery logs, seller performance records, and customer reviews. Extracting answers from this dataset "
        "('What was our top-selling product category in Rio de Janeiro last quarter?') requires advanced proficiency "
        "in Structured Query Language (SQL)."
    )
    add_bullet(
        doc,
        "2. Unstructured Operational Documents:",
        "Critical corporate rules, terms of service, warranty stipulations, refund guidelines, and shipping "
        "Service Level Agreements (SLAs) embedded within lengthy static documents, text files, and policy PDFs. "
        "Answering compliance or service questions ('Can a customer request a refund for opened electronics after 10 days?') "
        "demands laborious, manual document lookup by human agents."
    )

    format_heading_2(doc, "1.3 The Core Problem Statement")
    add_body_paragraph(
        doc,
        "Non-technical business stakeholders, customer support personnel, operations managers, and executive leadership "
        "face an acute information-access bottleneck. They cannot self-serve answers from either their operational "
        "databases or their corporate policy repositories without constant engineering dependency or time-consuming manual research. "
        "Traditional solutions only address one side of this divide—either offering a standalone Text-to-SQL tool or a simple document RAG search. "
        "There is an absence of unified, context-aware GenAI systems capable of automatically classifying user queries, dynamically "
        "dispatching them to the appropriate modality, executing safe queries, and delivering verifiable, grounded answers."
    )
    add_callout(
        doc,
        "The primary bottleneck is not a lack of organizational data, but the absence of an autonomous bridge between non-technical "
        "natural language questions and the disparate storage paradigms (relational tables vs. dense vector spaces) housing that data.",
        bold_prefix="Core Problem Synthesis: "
    )

    # ---------------------------------------------------------
    # 2. TARGET USERS & USE CASES
    # ---------------------------------------------------------
    format_heading_1(doc, "2. Target Users & Enterprise Use Cases")
    
    format_heading_2(doc, "2.1 Primary Target Stakeholders")
    add_bullet(
        doc,
        "Customer Support & Service Representatives:",
        "Frontline staff managing high-frequency customer inquiries. They currently toggle between database portals "
        "(to check order delivery dates and payment status) and PDF policy manuals (to verify return eligibility and compensation rules). "
        "The system provides immediate, contextual resolution in seconds."
    )
    add_bullet(
        doc,
        "Business & Operations Analysts:",
        "Mid-level analysts monitoring freight delays, logistics performance, seller metrics, and product revenue trends. "
        "Rather than filing ad-hoc SQL query requests with data engineering backlogs, analysts can interrogate the live "
        "PostgreSQL database directly using conversational natural language."
    )
    add_bullet(
        doc,
        "Executive Leadership & Department Heads:",
        "Decision-makers seeking rapid macro-level Key Performance Indicators (KPIs) during high-stakes reviews. "
        "They receive instantaneous, data-backed answers without waiting days for static dashboard alterations."
    )
    add_bullet(
        doc,
        "Compliance & Legal Officers:",
        "Personnel responsible for verifying operations against consumer protection legislation, vendor contracts, "
        "and dispute resolution guidelines. The system pinpoints exact policy clauses and provides verifiable section citations."
    )

    format_heading_2(doc, "2.2 Real-World Use Case Scenarios")
    
    format_heading_3(doc, "Scenario A: High-Velocity Customer Service Resolution")
    add_body_paragraph(
        doc,
        "A customer contacts support inquiring why their package has not arrived and whether they are entitled to an expedited shipping credit. "
        "A support representative enters: 'Order #38291 is 4 days past estimated delivery. What is our policy regarding late delivery compensation?' "
        "The system routes the question to the Document RAG engine, retrieves the exact logistics compensation clause, cites Section 3.4 of the "
        "Shipping SLA, and formats a customer-ready explanation."
    )

    format_heading_3(doc, "Scenario B: Zero-Code Ad-Hoc Business Intelligence")
    add_body_paragraph(
        doc,
        "An e-commerce category manager asks: 'Which product categories generated the top 10% of revenue in São Paulo with an average review score below 3.5?' "
        "The system routes to the SQL agent, inspects live schema relations across orders, products, reviews, and customer locations, generates "
        "a multi-table JOIN query, executes it safely, and presents both a visual summary and a structured data table."
    )

    # ---------------------------------------------------------
    # 3. WHAT YOUR GENAI SYSTEM WILL DO
    # ---------------------------------------------------------
    format_heading_1(doc, "3. System Architecture & Capabilities")
    add_body_paragraph(
        doc,
        "The GenAI system is architected as an autonomous multi-agent state machine orchestrated by LangGraph. "
        "Instead of treating an LLM as a passive conversational chatbot, the system establishes a proactive supervisor-subagent "
        "topology designed for deterministic execution, high accuracy, and strict security."
    )

    format_heading_2(doc, "3.1 Supervisor Router (The Agentic Brain)")
    add_body_paragraph(
        doc,
        "The central LangGraph orchestrator acts as the query dispatcher. It intercepts natural language prompts and conducts semantic "
        "intent classification using zero-shot reasoning. It classifies queries into three operational channels:"
    )
    add_bullet(doc, "• SQL Sub-Agent:", "Directed when queries involve operational metrics, calculations, counts, sums, rankings, or dates.")
    add_bullet(doc, "• RAG Sub-Agent:", "Directed when queries involve corporate rules, policies, terms of service, procedures, or SLAs.")
    add_bullet(doc, "• Conversational Channel:", "Handles greetings, clarifications, and capability explanations directly.")

    format_heading_2(doc, "3.2 Schema-Aware Text-to-SQL Generation Sub-Agent")
    add_body_paragraph(
        doc,
        "The SQL Sub-Agent dynamically interrogates PostgreSQL schema catalogs upon initialization. It understands the table "
        "topologies of 9 interconnected e-commerce entities (orders, customers, order_items, payments, reviews, products, sellers, etc.). "
        "When invoked, it performs:"
    )
    add_bullet(doc, "• Context Injection:", "Injects accurate table names, column data types, and primary-foreign key relationships into the prompt.")
    add_bullet(doc, "• Query Formulation:", "Synthesizes standard PostgreSQL queries incorporating complex JOINs, GROUP BY aggregations, and window functions.")
    add_bullet(doc, "• Dialect Validation:", "Ensures strict PostgreSQL syntax adherence (e.g., date arithmetic using INTERVAL and timestamp parsing).")

    format_heading_2(doc, "3.3 Defensive Security & Read-Only Query Guardrails")
    add_body_paragraph(
        doc,
        "To mitigate prompt injection attacks and protect database integrity, the Text-to-SQL engine features an automated AST and regex "
        "sanitization layer. Any query containing mutating or administrative commands—including DROP, DELETE, INSERT, UPDATE, ALTER, "
        "TRUNCATE, or GRANT—is blocked prior to execution, raising an explicit security alert."
    )

    format_heading_2(doc, "3.4 Grounded Document RAG & Anti-Hallucination Engine")
    add_body_paragraph(
        doc,
        "The Document RAG component enforces strict factual grounding over corporate documentation:"
    )
    add_bullet(doc, "• Semantic Ingestion & Chunking:", "Parses policy handbooks using RecursiveCharacterTextSplitter (500-character chunks, 50-character overlap) to preserve semantic coherence across clause boundaries.")
    add_bullet(doc, "• Vector Storage & Similarity Search:", "Converts chunks into 384-dimensional dense embeddings via a local sentence-transformer and builds a FAISS index for sub-millisecond retrieval.")
    add_bullet(doc, "• Citation Generation:", "Constrains the LLM to synthesize responses exclusively from retrieved chunks, mandating section-level citations (e.g., [Policy Section 2.1]).")

    format_heading_2(doc, "3.5 Dual-Surface Presentation Layer (Streamlit & CLI)")
    add_body_paragraph(
        doc,
        "Users interact through a polished Streamlit web application providing complete transparency: tabs for conversational answers, "
        "interactive Plotly visualizations, generated SQL query inspection panels, and source-chunk citation audits. "
        "A standalone terminal CLI loop is also provided for headless batch evaluation and automated testing."
    )
    add_figure(
        doc,
        "screenshot_dashboard.png",
        "Figure 3.1: Streamlit Prototype Dashboard showing system status telemetry, model configuration, sample prompt controls, and conversational interface.",
        width_inches=3.6
    )

    # ---------------------------------------------------------
    # 4. EXPECTED INPUT & OUTPUT
    # ---------------------------------------------------------
    format_heading_1(doc, "4. Expected Input & Output Specifications")
    
    format_heading_2(doc, "4.1 Input Characteristics")
    add_body_paragraph(
        doc,
        "The system accepts free-form, unformatted natural language text strings from users. The user is not required "
        "to specify routing flags, mention table names, or understand underlying data schemas."
    )

    format_heading_2(doc, "4.2 Multi-Layered Output Structure")
    add_body_paragraph(
        doc,
        "Depending on the routed execution path, the engine returns a multi-layered response payload designed for both "
        "immediate human comprehension and developer auditing:"
    )
    add_bullet(doc, "1. Conversational Executive Summary:", "A synthesized natural language response formatting metrics (currency, percentages, timelines) into clear business English.")
    add_bullet(doc, "2. Transparent Engineering Artifacts:", "The raw executed SQL query string, data preview tables, or retrieved policy text chunks with similarity rankings.")
    add_bullet(doc, "3. System Telemetry Metadata:", "Routing decision tag, database query execution latency (seconds), and total tokens consumed.")

    format_heading_2(doc, "4.3 Comprehensive Input-to-Output Mapping Matrix")
    
    headers_io = ["User Query (Input)", "Routing Path", "Synthesized Response (Primary Output)", "Auditing Artifacts"]
    data_io = [
        [
            "What are the top 5 product categories by revenue in São Paulo?",
            "SQL Sub-Agent\n(PostgreSQL)",
            "Identifies top categories with exact revenue totals, lead counts, and average unit pricing in business narrative.",
            "Generated PostgreSQL query, Pandas preview table (5 rows), query latency."
        ],
        [
            "Can a customer return an opened electronic item after 10 days?",
            "RAG Sub-Agent\n(FAISS Vector Store)",
            "Explains that standard returns are 7 days for opened electronics, citing warranty inspection rules and refund deduction policies.",
            "Retrieved chunks with similarity scores, Policy Section 4.2 cited."
        ],
        [
            "Which seller had the fastest average delivery time in 2017?",
            "SQL Sub-Agent\n(PostgreSQL)",
            "Calculates order delivery timestamp minus purchase timestamp, identifying the top seller with an average of 4.2 days.",
            "SQL query with timestamp interval math, tabular seller ranking."
        ],
        [
            "DROP TABLE customers;",
            "Safety Guardrail\n(Security Validator)",
            "Execution Denied: 'Destructive DDL/DML operations are strictly blocked. Only read-only SELECT queries are permitted.'",
            "Security rule violation log, AST rejection flag."
        ]
    ]
    create_table(doc, headers_io, data_io, col_widths=[1.8, 1.2, 2.0, 1.5])

    format_heading_2(doc, "4.4 Prototype Execution Evidence & Screen Demonstrations")
    add_body_paragraph(
        doc,
        "The figures below demonstrate the live execution of the dual-path agentic engine against the PostgreSQL relational database "
        "and the FAISS vector store in our Streamlit prototype:"
    )

    format_heading_3(doc, "4.4.1 Text-to-SQL Relational Retrieval in Action")
    add_body_paragraph(
        doc,
        "When prompted with an operational inquiry ('What are the top 5 product categories by number of orders?'), the supervisor routes "
        "the request to the SQL Data Agent. The agent inspects the relational schema, crafts and executes a multi-table aggregation query, "
        "and returns a structured ranking table with exact order frequencies:"
    )
    add_figure(
        doc,
        "screenshot_sql.png",
        "Figure 4.1: Live Text-to-SQL Execution demonstrating prompt routing, SQL Data Agent badge, and structured order frequency breakdown.",
        width_inches=6.0
    )

    format_heading_3(doc, "4.4.2 Document RAG Policy Retrieval with Citations")
    add_body_paragraph(
        doc,
        "When prompted with a customer compliance inquiry ('What is the policy for returning a defective or damaged item?'), the router directs "
        "the query to the Policy RAG Engine. The engine retrieves relevant policy passages from FAISS and synthesizes an authoritative answer "
        "with exact section citations:"
    )
    add_figure(
        doc,
        "screenshot_rag.png",
        "Figure 4.2: Live Document RAG Engine demonstrating policy query classification, grounded bullet-point resolution, and exact clause citation (Section 2: Defective or Damaged Goods).",
        width_inches=6.0
    )

    # ---------------------------------------------------------
    # 5. TOOLS & MODELS PLAN
    # ---------------------------------------------------------
    format_heading_1(doc, "5. Tools, Frameworks & Model Architecture")
    add_body_paragraph(
        doc,
        "The solution is constructed upon an enterprise-ready open-source and API foundation, balancing state-of-the-art "
        "reasoning capability with local privacy and execution speed."
    )

    headers_tools = ["Component Layer", "Tool / Technology", "Specification / Version", "Architectural Justification"]
    data_tools = [
        [
            "Generative Reasoning",
            "OpenAI GPT-4o-mini\n(Groq Llama-3 Fallback)",
            "Temperature = 0.0\nTop-P = 1.0",
            "Optimal balance of schema reasoning, strict SQL dialect adherence, low latency, and deterministic output."
        ],
        [
            "Dense Vector Embeddings",
            "Sentence-Transformers\nall-MiniLM-L6-v2",
            "384 Dimensions\nLocal CPU inference",
            "Runs entirely on-premise/locally with zero token cost and microsecond vector generation; proven semantic search accuracy."
        ],
        [
            "Agent Orchestration",
            "LangGraph & LangChain",
            "LangGraph >= 1.2.0\nLangChain >= 1.3.0",
            "Enables cyclic state graphs, deterministic conditional routing, error recovery loops, and clean modular sub-agents."
        ],
        [
            "Relational Database",
            "PostgreSQL & SQLAlchemy",
            "PostgreSQL 15+\nSQLAlchemy >= 2.0.0",
            "ACID-compliant relational engine holding 100K+ e-commerce orders across 9 interrelated tables."
        ],
        [
            "Vector Database",
            "FAISS (faiss-cpu)",
            "IndexFlatL2 / Cosine\nIn-memory indexing",
            "Lightning-fast dense vector similarity search without requiring heavy external cloud vector database infrastructure."
        ],
        [
            "User Interface & Viz",
            "Streamlit & Plotly",
            "Streamlit >= 1.30.0\nPlotly >= 5.18.0",
            "Rapid, reactive web interface supporting interactive data exploration, chart rendering, and audit toggles."
        ],
        [
            "Verification & Quality",
            "Pytest & LangSmith",
            "Pytest >= 7.4.0\nLangSmith Observability",
            "Automated unit testing for SQL sanitization and agent evaluation; end-to-end token tracing."
        ]
    ]
    create_table(doc, headers_tools, data_tools, col_widths=[1.4, 1.6, 1.4, 2.1])

    format_heading_2(doc, "5.1 Automated System Verification & Test Suite Results")
    add_body_paragraph(
        doc,
        "To guarantee execution safety and deterministic routing, the codebase incorporates a comprehensive automated test suite "
        "implemented in Pytest. The suite covers SQL mutation gate validation (blocking dangerous DDL/DML), vector chunk retrieval accuracy, "
        "and LangGraph state transitions:"
    )
    add_figure(
        doc,
        "screenshot_tests.png",
        "Figure 5.1: Automated Verification Test Suite (Pytest) showing 100% pass rate (7/7 tests passed) across SQL security gates, vector chunking, and LangGraph routing.",
        width_inches=6.0
    )

    # ---------------------------------------------------------
    # 6. ONE THING YOU ARE CURRENTLY UNSURE ABOUT
    # ---------------------------------------------------------
    format_heading_1(doc, "6. Technical Uncertainty & Research Direction")
    
    format_heading_2(doc, "6.1 The Core Challenge: Compound Cross-Modal (Hybrid) Queries")
    add_body_paragraph(
        doc,
        "The primary architectural uncertainty currently under active investigation is the robust handling of "
        "compound 'cross-modal' questions that simultaneously require both unstructured document policy knowledge "
        "and structured relational database calculations in a single prompt."
    )
    add_body_paragraph(
        doc,
        "Our current supervisor router functions on a discrete, either-or classification paradigm: a query is directed "
        "exclusively to either the SQL sub-agent or the Document RAG engine. However, real-world enterprise operations "
        "frequently generate intertwined questions."
    )

    format_heading_3(doc, "Concrete Problem Illustration:")
    add_callout(
        doc,
        "User Prompt: 'Identify all sellers who violated our 5-day dispatch SLA during the Q4 promotional campaign.'\n\n"
        "• The Relational Database contains order timestamps, seller IDs, and shipping dates, but has no intrinsic definition of what the 'dispatch SLA' stipulates.\n"
        "• The Policy Document defines the exact SLA threshold (5 business days), but has zero awareness of seller orders or transaction logs.\n"
        "• Under the current either-or routing model, routing to SQL fails due to an undefined SLA threshold, while routing to RAG yields policy text without seller data.",
        bold_prefix="The Cross-Modal Dilemma: "
    )

    format_heading_2(doc, "6.2 Architectural Options Under Evaluation")
    add_body_paragraph(
        doc,
        "To resolve this cross-modal dependency without introducing fragile brittle logic, we are evaluating three distinct architectural patterns:"
    )
    add_bullet(
        doc,
        "1. Dynamic Sub-Goal Planning (DAG Decomposition):",
        "Empowering the supervisor router to detect cross-modal intent and compile a Directed Acyclic Graph (DAG) plan. "
        "Step 1 queries the RAG engine to extract the parameter value ('SLA = 5 days'), and Step 2 injects that extracted "
        "parameter dynamically into the SQL agent's generation prompt as a WHERE clause filter."
    )
    add_bullet(
        doc,
        "2. Iterative Tool-Calling with Shared Working Memory:",
        "Transitioning from a static router to a ReAct (Reasoning + Acting) loop where a single master agent can call the "
        "RAG tool, inspect the retrieved document context, store intermediate facts in scratchpad memory, and subsequently "
        "call the SQL tool to execute the final query."
    )
    add_bullet(
        doc,
        "3. Semantic SQL Logic Verification (Correctness vs. Validity):",
        "Addressing the subtle challenge of ensuring the LLM selects the correct business logic column (e.g., choosing "
        "between order_purchase_timestamp and order_delivered_carrier_date) when business colloquialisms like 'delayed order' "
        "are used. We are experimenting with few-shot schema semantic annotations and an automated SQL reflection reviewer."
    )

    # Save Document
    output_filename = "Project_Report_Agentic_Text_to_SQL_and_RAG_Engine.docx"
    output_path = os.path.join(os.path.dirname(__file__), output_filename)
    doc.save(output_path)
    print(f"Report successfully saved to: {output_path}")

if __name__ == "__main__":
    generate_report()
